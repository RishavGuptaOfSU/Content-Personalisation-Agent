#!/usr/bin/env python
"""Checks attachment extraction and the vision plumbing — no network.

Covers the three things that can go quietly wrong:

* a PDF or DOCX that extracts to nothing (a scan, an encrypted file) must raise
  a message the user can act on, not return an empty attachment that makes the
  model look like it ignored the document
* an image must reach the provider as image data, in each provider's own wire
  format, and never be silently dropped
* a model that cannot see must not be handed an image at all

    python scripts/check_attachments.py
"""

from __future__ import annotations

import base64
import io
import sys
import zlib
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

FAILURES: list[str] = []


def check(label: str, condition: bool, detail: str = "") -> None:
    mark = "PASS" if condition else "FAIL"
    print(f"  [{mark}] {label}" + (f" -> {detail}" if detail else ""))
    if not condition:
        FAILURES.append(label)


# ------------------------------------------------------------------ fixtures
def make_pdf(pages: list[str]) -> bytes:
    """A minimal multi-page PDF with a real text layer, built by hand.

    Written rather than fixtured so the test has no binary blobs in the repo
    and the expected text is visible right here.
    """
    objects: list[bytes] = []
    page_ids = [4 + index * 2 for index in range(len(pages))]

    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(
        f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode()
    )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    for index, text in enumerate(pages):
        content = (
            f"BT /F1 24 Tf 72 700 Td ({text}) Tj ET".encode()
        )
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 3 0 R >> >> "
                f"/Contents {page_ids[index] + 1} 0 R >>"
            ).encode()
        )
        objects.append(
            b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n"
            + content + b"\nendstream"
        )

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF\n"
    ).encode()
    return bytes(out)


def make_image_pdf() -> bytes:
    """A PDF with a drawing and no text layer — i.e. what a scan looks like."""
    content = b"0 0 1 rg 100 100 300 300 re f"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF\n"
    ).encode()
    return bytes(out)


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for col_index, cell in enumerate(row):
                grid.cell(row_index, col_index).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_png(width: int = 60, height: int = 40) -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), (40, 70, 200)).save(buffer, format="PNG")
    return buffer.getvalue()


def main() -> int:
    from app.services import attachments
    from app.services.attachments import AttachmentError, extract

    print("1. plain text")
    result = extract("notes.txt", b"Acme Cloud Q3 plan\nbudget 400 USD")
    check("kind is text", result.kind == "text", result.kind)
    check("content is preserved", "Acme Cloud Q3 plan" in result.text)
    check("no data url for text", result.data_url == "")

    bad_bytes = b"caf\xe9 latte"  # valid latin-1, invalid utf-8
    check(
        "non-utf8 text still decodes",
        "caf" in extract("notes.txt", bad_bytes).text,
    )

    print("\n2. PDF with a text layer")
    pdf = make_pdf(["Acme Cloud rightsizing brief", "Budget is 400 USD"])
    result = extract("brief.pdf", pdf)
    check("kind is text", result.kind == "text", result.kind)
    check(
        "page one text extracted",
        "Acme Cloud rightsizing brief" in result.text,
        result.text[:60].replace("\n", " "),
    )
    check("page two text extracted", "Budget is 400 USD" in result.text)
    check("it reports the page count", "2 PDF page" in " ".join(result.notes), str(result.notes))

    print("\n3. a PDF with no text layer is refused, not returned empty")
    try:
        extract("scan.pdf", make_image_pdf())
    except AttachmentError as exc:
        message = str(exc)
        check("it raises", True)
        check("it says no text was found", "No text found" in message)
        check("it explains OCR is not done", "OCR" in message)
        check("it suggests what to do instead", "PNG" in message or "Paste" in message)
    else:
        check("it raises", False, "returned an attachment for a scan")

    print("\n4. a corrupt PDF fails with a readable message")
    for label, blob in {
        "random bytes": b"\x00\x01\x02not a pdf at all",
        "truncated pdf": make_pdf(["hello"])[:40],
    }.items():
        try:
            extract("broken.pdf", blob)
        except AttachmentError:
            check(f"rejects {label}", True)
        except Exception as exc:  # noqa: BLE001
            check(f"rejects {label}", False, f"{type(exc).__name__} escaped: {exc}")
        else:
            check(f"rejects {label}", False, "parsed without complaint")

    print("\n5. DOCX, including tables")
    docx_bytes = make_docx(
        ["Launch brief", "Audience: developers and founders"],
        table=[["Channel", "Budget"], ["LinkedIn", "200"], ["Email", "150"]],
    )
    result = extract("brief.docx", docx_bytes)
    check("kind is text", result.kind == "text")
    check("paragraphs extracted", "Launch brief" in result.text)
    check("second paragraph extracted", "developers and founders" in result.text)
    check(
        "table cells extracted",
        "LinkedIn" in result.text and "200" in result.text,
        "tables are where briefs keep their numbers",
    )
    check("it reports the table", "table" in " ".join(result.notes), str(result.notes))

    try:
        extract("legacy.doc", b"\xd0\xcf\x11\xe0old word format")
    except AttachmentError as exc:
        check("a .doc is refused by suffix", "not supported" in str(exc).lower())
    else:
        check("a .doc is refused by suffix", False)

    try:
        extract("broken.docx", b"PK\x03\x04 but not really a docx")
    except AttachmentError as exc:
        check("a non-docx .docx is refused", ".docx" in str(exc), str(exc)[:72])
    else:
        check("a non-docx .docx is refused", False)

    print("\n6. empty and oversized content")
    try:
        extract("empty.txt", b"   \n  ")
    except AttachmentError as exc:
        check("whitespace-only file is refused", "empty" in str(exc).lower())
    else:
        check("whitespace-only file is refused", False)

    long_text = ("x" * 100 + "\n") * 600  # ~60k chars
    result = extract("big.txt", long_text.encode())
    check(
        "long text is truncated to the cap",
        result.characters == attachments.MAX_TEXT_CHARS,
        f"{result.characters} chars",
    )
    check("truncation is flagged", result.truncated is True)
    check("truncation is explained", any("Truncated" in n for n in result.notes))

    print("\n7. images become vision input, not text")
    png = make_png(60, 40)
    result = extract("shot.png", png)
    check("kind is image", result.kind == "image", result.kind)
    check("a data url is produced", result.data_url.startswith("data:image/png;base64,"))
    check("dimensions are read", (result.width, result.height) == (60, 40),
          f"{result.width}x{result.height}")
    check("no text is invented", result.text == "")
    check(
        "the payload round-trips to the original bytes",
        base64.b64decode(result.data_url.split(",", 1)[1]) == png,
    )
    check("jpeg is accepted", extract("a.jpg", _jpeg()).kind == "image")

    try:
        extract("fake.png", b"this is not a png")
    except AttachmentError as exc:
        check("a fake image is refused", "not a readable image" in str(exc), str(exc)[:50])
    else:
        check("a fake image is refused", False)

    try:
        extract("archive.zip", b"PK\x03\x04")
    except AttachmentError as exc:
        check("an unsupported suffix is refused", "not supported" in str(exc))
    else:
        check("an unsupported suffix is refused", False)

    print("\n8. vision capability is decided from the model name")
    from app.core.config import settings
    from app.llm import vision

    original_provider, original_setting = settings.LLM_PROVIDER, settings.LLM_VISION
    original_model = settings.OPENAI_MODEL
    try:
        settings.LLM_PROVIDER, settings.LLM_VISION = "openai", "auto"
        for model, expected in {
            "gemma4:31b": True,          # measured: reads images
            "gpt-oss:120b": False,       # measured: "does not support image input"
            "gpt-oss:20b": False,
            "beam": False,
            "nemotron-3-super": False,
            "nemotron-3-ultra": False,
            "gpt-4o-mini": True,
            "claude-sonnet-4": True,
            "llama3.2-vision:11b": True,
            "llama3.2:3b": False,
        }.items():
            settings.OPENAI_MODEL = model
            check(f"{model} vision={expected}", vision.vision_supported() is expected)

        settings.OPENAI_MODEL = "some-new-vision-model"
        settings.LLM_VISION = "on"
        check("LLM_VISION=on overrides detection", vision.vision_supported() is True)
        settings.OPENAI_MODEL = "gemma4:31b"
        settings.LLM_VISION = "off"
        check("LLM_VISION=off refuses a capable model", vision.vision_supported() is False)
        settings.LLM_VISION = "auto"
        status = vision.vision_status()
        check("status reports the model", status["model"] == "gemma4:31b", str(status["model"]))
        check("status explains itself", bool(status["reason"]))
    finally:
        settings.LLM_PROVIDER = original_provider
        settings.LLM_VISION = original_setting
        settings.OPENAI_MODEL = original_model

    print("\n9. an image reaches each provider in its own format")
    from app.llm.base import ChatMessage
    from app.llm.ollama import _normalize_messages

    data_url = f"data:image/png;base64,{base64.b64encode(png).decode()}"
    message = ChatMessage(role="user", content="What does this say?", images=[data_url])

    openai_shape = message.to_dict()
    check("openai content becomes a parts list", isinstance(openai_shape["content"], list))
    kinds = [part["type"] for part in openai_shape["content"]]
    check("the text part comes first", kinds == ["text", "image_url"], str(kinds))
    check(
        "the image part carries the data url",
        openai_shape["content"][1]["image_url"]["url"] == data_url,
    )
    check(
        "a text-only message stays a plain string",
        isinstance(ChatMessage(role="user", content="hi").to_dict()["content"], str),
    )

    ollama_shape = _normalize_messages([message])[0]
    check("ollama keeps content as a string", isinstance(ollama_shape["content"], str))
    check("ollama uses a sibling images list", "images" in ollama_shape)
    check(
        "ollama gets bare base64, not a data url",
        not ollama_shape["images"][0].startswith("data:"),
        ollama_shape["images"][0][:24],
    )
    check(
        "the bare base64 still decodes to the image",
        base64.b64decode(ollama_shape["images"][0]) == png,
    )

    # Merging exists to satisfy strict chat templates, but it must never fold an
    # image into a neighbouring turn.
    merged = _normalize_messages([
        ChatMessage(role="user", content="context block"),
        message,
    ])
    check("an image turn is not merged away", len(merged) == 2, f"{len(merged)} messages")
    check("the image survives merging", "images" in merged[-1])
    plain_merged = _normalize_messages([
        ChatMessage(role="user", content="one"),
        ChatMessage(role="user", content="two"),
    ])
    check("text turns still merge", len(plain_merged) == 1, f"{len(plain_merged)} messages")

    print("\n10. a provider that cannot send images refuses loudly")
    from app.llm.base import LLMProviderError
    from app.llm.bedrock import BedrockLLMProvider

    try:
        BedrockLLMProvider.generate(
            object.__new__(BedrockLLMProvider), [message], system=None
        )
    except LLMProviderError as exc:
        check("bedrock raises rather than dropping the image", "not implemented" in str(exc))
        check("it names a way forward", "LLM_PROVIDER=openai" in str(exc))
    except Exception as exc:  # noqa: BLE001
        check(
            "bedrock raises rather than dropping the image",
            False,
            f"{type(exc).__name__}: {exc}",
        )
    else:
        check("bedrock raises rather than dropping the image", False, "accepted it")

    print("\n11. the context carries images to the final user turn")
    from app.agents.catalog import get_agent_spec
    from app.services.personalization import PersonalizedContext

    context = PersonalizedContext(
        agent_key="marketing.social-copy",
        agent_spec=get_agent_spec("marketing.social-copy"),
        system_prompt="",
        context_block="profile stuff",
        short_term=[],
        memories=[],
        hints={},
        user_request="Write a caption for this",
        attachment_images=[data_url],
    )
    messages = context.to_messages()
    check("the last message is the user request", messages[-1].role == "user")
    check("it carries the image", messages[-1].images == [data_url])
    check(
        "the context block does not carry the image",
        all(not m.images for m in messages[:-1]),
        "the model is told to ignore that block",
    )
    check(
        "a text-only context carries no images",
        not any(
            m.images
            for m in PersonalizedContext(
                agent_key="marketing.social-copy",
                agent_spec=get_agent_spec("marketing.social-copy"),
                system_prompt="",
                context_block="",
                short_term=[],
                memories=[],
                hints={},
                user_request="hello",
            ).to_messages()
        ),
    )

    print()
    if FAILURES:
        print(f"FAILED ({len(FAILURES)}): " + "; ".join(FAILURES))
        return 1
    print("ALL ATTACHMENT CHECKS PASSED")
    return 0


def _jpeg() -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (32, 32), (200, 40, 40)).save(buffer, format="JPEG")
    return buffer.getvalue()


if __name__ == "__main__":
    del zlib  # imported for the PDF builder's earlier draft; keep the lint clean
    raise SystemExit(main())
