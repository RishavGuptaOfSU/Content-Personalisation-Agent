"""Turning an uploaded file into something a model can actually use.

Two shapes come out of here, and the difference matters:

* **text** — a ``.txt``, ``.csv``, ``.pdf`` or ``.docx`` becomes prose that is
  appended to the prompt. Any model can read it.
* **image** — a ``.png`` or ``.jpg`` stays binary and is handed to the model as
  vision input. Only a multimodal model can read it, so the caller has to check
  before promising the user anything.

Why extraction is server-side: the browser would have to ship a PDF parser to
do it client-side, and the raw bytes of a PDF are useless to a text model —
latin-1 will happily "decode" them into line noise, which is worse than an
error because it looks like it worked.

No OCR. A scanned PDF has no text layer, and silently returning an empty
attachment would leave the user wondering why the model ignored their document,
so that case raises instead.
"""

from __future__ import annotations

import base64
import io
from dataclasses import dataclass, field
from typing import Literal

from app.core.logging import get_logger

logger = get_logger(__name__)

#: Upper bound on extracted characters handed to the prompt. Matches the
#: ``attachment_text`` ceiling in the chat schema.
MAX_TEXT_CHARS = 40_000

#: Suffixes that become prompt text, grouped by how they are read.
PLAIN_TEXT_SUFFIXES = (
    ".txt", ".md", ".csv", ".json", ".tsv", ".log", ".py", ".sql", ".yaml", ".yml",
)
PDF_SUFFIXES = (".pdf",)
DOCX_SUFFIXES = (".docx",)

#: Suffixes handed to the model as vision input, with their media types.
IMAGE_MEDIA_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
}

TEXT_SUFFIXES = PLAIN_TEXT_SUFFIXES + PDF_SUFFIXES + DOCX_SUFFIXES
SUPPORTED_SUFFIXES = TEXT_SUFFIXES + tuple(IMAGE_MEDIA_TYPES)

AttachmentKind = Literal["text", "image"]


class AttachmentError(ValueError):
    """Raised with a message intended for the user, not the log."""


@dataclass
class Attachment:
    """What the upload endpoint hands back to the client."""

    name: str
    kind: AttachmentKind
    #: Extracted prose. Empty for images.
    text: str = ""
    truncated: bool = False
    #: ``data:image/png;base64,...`` for images. Empty for text.
    data_url: str = ""
    width: int | None = None
    height: int | None = None
    #: Human-readable note about how the content was obtained.
    notes: list[str] = field(default_factory=list)

    @property
    def characters(self) -> int:
        return len(self.text)

    def to_dict(self) -> dict[str, object]:
        payload: dict[str, object] = {
            "name": self.name,
            "kind": self.kind,
            "characters": self.characters,
            "truncated": self.truncated,
            "text": self.text,
            "notes": self.notes,
        }
        if self.kind == "image":
            payload |= {
                "data_url": self.data_url,
                "width": self.width,
                "height": self.height,
            }
        return payload


# --------------------------------------------------------------------- utils
def suffix_of(filename: str) -> str:
    name = (filename or "").lower()
    dot = name.rfind(".")
    return name[dot:] if dot != -1 else ""


def is_image(filename: str) -> bool:
    return suffix_of(filename) in IMAGE_MEDIA_TYPES


def _truncate(text: str) -> tuple[str, bool]:
    cleaned = text.replace("\x00", "").strip()
    if len(cleaned) <= MAX_TEXT_CHARS:
        return cleaned, False
    return cleaned[:MAX_TEXT_CHARS], True


# ------------------------------------------------------------------ readers
def _read_plain_text(raw: bytes) -> str:
    """Decode a text file, preferring utf-8.

    latin-1 is the fallback because it cannot fail — every byte maps to a
    character. That makes it a safety net for odd encodings, not a validity
    check, which is why the suffix allowlist upstream is what actually keeps
    binary files out of this path.
    """
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        logger.debug("Attachment is not utf-8; falling back to latin-1")
        return raw.decode("latin-1")


def _read_pdf(raw: bytes) -> tuple[str, list[str]]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(raw))
    except PdfReadError as exc:
        raise AttachmentError(f"This PDF could not be read ({exc}).") from exc

    if reader.is_encrypted:
        # An empty password unlocks a surprising number of "protected" PDFs.
        try:
            if reader.decrypt("") == 0:
                raise AttachmentError(
                    "This PDF is password protected. Remove the password and "
                    "upload it again."
                )
        except AttachmentError:
            raise
        except Exception as exc:  # noqa: BLE001 - pypdf raises various types here
            raise AttachmentError(
                "This PDF is encrypted in a way we cannot open."
            ) from exc

    pages: list[str] = []
    failed = 0
    for index, page in enumerate(reader.pages, start=1):
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - one broken page should not lose the rest
            failed += 1
            logger.warning("Could not extract text from PDF page %d", index)
            pages.append("")

    text = "\n\n".join(part.strip() for part in pages if part.strip())
    if not text.strip():
        raise AttachmentError(
            f"No text found in this PDF ({len(reader.pages)} page(s)). It is "
            "probably a scan or an image-only export — this reads a PDF's text "
            "layer and does not do OCR. Paste the text, or upload the page as "
            "a PNG so the model can look at it instead."
        )

    notes = [f"Extracted the text layer of {len(reader.pages)} PDF page(s)."]
    if failed:
        notes.append(f"{failed} page(s) could not be read and were skipped.")
    return text, notes


def _read_docx(raw: bytes) -> tuple[str, list[str]]:
    import zipfile

    import docx
    from docx.opc.exceptions import OpcError, PackageNotFoundError

    # A .docx is a zip. python-docx surfaces a bad one as PackageNotFoundError,
    # but a file that is not a zip at all escapes as zipfile.BadZipFile from
    # deeper down, so both have to be caught to give one clear message.
    try:
        document = docx.Document(io.BytesIO(raw))
    except (PackageNotFoundError, OpcError, zipfile.BadZipFile, KeyError) as exc:
        raise AttachmentError(
            "This does not look like a readable .docx file. The old .doc format "
            "is not supported — open it and re-save as .docx."
        ) from exc

    blocks = [para.text.strip() for para in document.paragraphs if para.text.strip()]

    # Tables carry a lot of the useful content in briefs and specs, and
    # document.paragraphs skips them entirely.
    tables = 0
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            tables += 1
            blocks.append("\n".join(rows))

    text = "\n\n".join(blocks)
    if not text.strip():
        raise AttachmentError(
            "This .docx has no readable text. If the content is all images, "
            "upload those as PNGs instead."
        )

    notes = [f"Read {len(document.paragraphs)} paragraph(s) from the .docx."]
    if tables:
        notes.append(f"Included {tables} table(s).")
    return text, notes


def _read_image(raw: bytes, suffix: str) -> Attachment:
    from PIL import Image, UnidentifiedImageError

    try:
        with Image.open(io.BytesIO(raw)) as image:
            image.verify()  # cheap integrity check; invalidates the handle
        with Image.open(io.BytesIO(raw)) as image:
            width, height = image.size
            fmt = (image.format or "").upper()
    except (UnidentifiedImageError, OSError) as exc:
        raise AttachmentError(
            "This file is not a readable image, or it is corrupt."
        ) from exc

    media_type = IMAGE_MEDIA_TYPES[suffix]
    encoded = base64.b64encode(raw).decode("ascii")
    return Attachment(
        name="",
        kind="image",
        data_url=f"data:{media_type};base64,{encoded}",
        width=width,
        height=height,
        notes=[f"{fmt or media_type} image, {width}x{height}, sent to the model to look at."],
    )


# -------------------------------------------------------------------- entry
def extract(filename: str, raw: bytes) -> Attachment:
    """Read an uploaded file into prompt text or vision input.

    Raises ``AttachmentError`` with a message safe to show the user.
    """
    suffix = suffix_of(filename)
    display_name = filename or "attachment"

    if suffix in IMAGE_MEDIA_TYPES:
        attachment = _read_image(raw, suffix)
        attachment.name = display_name
        return attachment

    if suffix in PDF_SUFFIXES:
        text, notes = _read_pdf(raw)
    elif suffix in DOCX_SUFFIXES:
        text, notes = _read_docx(raw)
    elif suffix in PLAIN_TEXT_SUFFIXES:
        text, notes = _read_plain_text(raw), []
    else:
        raise AttachmentError(
            f"{suffix or 'This file type'} is not supported. "
            f"Supported: {', '.join(SUPPORTED_SUFFIXES)}"
        )

    truncated_text, truncated = _truncate(text)
    if not truncated_text:
        raise AttachmentError("That file is empty.")
    if truncated:
        notes.append(f"Truncated to the first {MAX_TEXT_CHARS:,} characters.")

    return Attachment(
        name=display_name,
        kind="text",
        text=truncated_text,
        truncated=truncated,
        notes=notes,
    )


__all__ = [
    "IMAGE_MEDIA_TYPES",
    "MAX_TEXT_CHARS",
    "SUPPORTED_SUFFIXES",
    "TEXT_SUFFIXES",
    "Attachment",
    "AttachmentError",
    "extract",
    "is_image",
    "suffix_of",
]
