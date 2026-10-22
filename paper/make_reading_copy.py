#!/usr/bin/env python
"""Renders a DOCX reading copy of the paper for review and sharing.

This is NOT the submission artefact — IEEE wants the LaTeX source and PDF. It
exists so the paper can be read, circulated to a supervisor and commented on
without a TeX installation. Tables are summarised rather than reproduced,
because faithfully rendering LaTeX tabular into DOCX is not worth the code.

    python3 paper/make_reading_copy.py
    soffice --headless --convert-to pdf --outdir paper paper/reading_copy.docx
"""

from __future__ import annotations

import re
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
TEX = HERE / "ieee_access_paper.tex"
OUT = HERE / "reading_copy.docx"

FIG_CAPTION_WIDTH = Inches(6.0)


def strip_latex(text: str) -> str:
    """Reduce LaTeX markup to readable prose."""
    text = re.sub(r"%.*", "", text)
    text = re.sub(r"\\IEEEPARstart\{(.)\}\{([^}]*)\}", r"\1\2", text)
    text = re.sub(r"\\(emph|textbf|textit|texttt)\{([^{}]*)\}", r"\2", text)
    text = re.sub(r"\\REPLACE\{([^}]*)\}", r"[REPLACE: \1]", text)
    text = re.sub(r"\\(ref|cite)\{([^}]*)\}", lambda m: (
        "[" + m.group(2).replace("_", " ") + "]"), text)
    text = re.sub(r"~", " ", text)
    text = text.replace(r"\,", " ").replace(r"\%", "%").replace(r"\$", "$")
    text = text.replace("---", "\u2014").replace("--", "\u2013")
    text = text.replace("``", "\u201c").replace("''", "\u201d")
    text = re.sub(r"\\[a-zA-Z]+\*?(\[[^\]]*\])?", " ", text)
    text = re.sub(r"[{}$\\]", "", text)
    return re.sub(r"[ \t]+", " ", text).strip()


def main() -> int:
    src = TEX.read_text(encoding="utf-8")
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(10.5)

    title = doc.add_heading(
        "Deterministic Guardrails and Variance-Aware Evaluation for "
        "Personalized Multi-Agent Content Generation", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph(
        "Rishav Gupta \u2014 Department of Computer Science and Engineering, "
        "Sharda University, Greater Noida, India")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note = doc.add_paragraph(
        "Reading copy. The submission artefact is ieee_access_paper.tex "
        "plus figures/. Target venue: IEEE Access.")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.runs[0].italic = True

    # Abstract
    abstract = src.split(r"\begin{abstract}")[1].split(r"\end{abstract}")[0]
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(strip_latex(abstract))

    keywords = src.split(r"\begin{IEEEkeywords}")[1].split(r"\end{IEEEkeywords}")[0]
    kw = doc.add_paragraph("Index Terms\u2014" + strip_latex(keywords))
    kw.runs[0].italic = True

    figures = {
        "fig1_pipeline": "figures/fig1_pipeline.png",
        "fig2_scale_collapse": "figures/fig2_scale_collapse.png",
        "fig3_bakeoff": "figures/fig3_bakeoff.png",
        "fig4_ab_per_agent": "figures/fig4_ab_per_agent.png",
    }

    body = src.split(r"\end{IEEEkeywords}")[1].split(r"\begin{thebibliography}")[0]

    # Walk the body, emitting headings, paragraphs, figures and table notices.
    token = re.compile(
        r"\\(section|subsection)\{([^}]*)\}"
        r"|\\begin\{(figure\*?|table)\}(.*?)\\end\{\3\}",
        re.DOTALL,
    )
    pos = 0
    for match in token.finditer(body):
        chunk = body[pos:match.start()]
        for para in re.split(r"\n\s*\n", chunk):
            text = strip_latex(para)
            if text:
                doc.add_paragraph(text)
        pos = match.end()

        if match.group(1):
            doc.add_heading(strip_latex(match.group(2)),
                            level=2 if match.group(1) == "section" else 3)
        else:
            env, content = match.group(3), match.group(4)
            caption = re.search(r"\\caption\{(.*?)\}\s*\\label", content, re.DOTALL)
            caption_text = strip_latex(caption.group(1)) if caption else ""
            label = re.search(r"\\label\{([^}]*)\}", content)
            key = label.group(1).split(":")[-1] if label else ""
            if env.startswith("figure"):
                name = re.search(r"\\includegraphics\[[^\]]*\]\{([^}]*)\}", content)
                stem = Path(name.group(1)).stem if name else ""
                path = HERE / figures.get(stem, "")
                if path.is_file():
                    doc.add_picture(str(path), width=FIG_CAPTION_WIDTH)
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph(f"Figure \u2014 {caption_text}")
            else:
                cap = doc.add_paragraph(
                    f"[Table \u2014 {caption_text} "
                    f"(see the LaTeX source, label {key}, for the full table.)]")
            cap.runs[0].italic = True
            cap.runs[0].font.size = Pt(9)

    for para in re.split(r"\n\s*\n", body[pos:]):
        text = strip_latex(para)
        if text:
            doc.add_paragraph(text)

    # References
    doc.add_heading("References", level=2)
    bib = src.split(r"\begin{thebibliography}")[1].split(r"\end{thebibliography}")[0]
    for index, item in enumerate(re.split(r"\\bibitem\{[^}]*\}", bib)[1:], start=1):
        text = strip_latex(item)
        if text:
            para = doc.add_paragraph(f"[{index}] {text}")
            para.paragraph_format.space_after = Pt(3)
            para.runs[0].font.size = Pt(9)

    doc.save(OUT)
    print(f"wrote {OUT.relative_to(HERE.parent)}")
    print("convert to PDF with:")
    print("  soffice --headless --convert-to pdf --outdir paper paper/reading_copy.docx")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
